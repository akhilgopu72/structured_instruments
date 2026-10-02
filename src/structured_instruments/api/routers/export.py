"""Excel export API router."""
from __future__ import annotations

import io
from datetime import date

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select, func
from sqlalchemy.orm import Session

from structured_instruments.api.dependencies import get_db, get_account_filter
from structured_instruments.models.coupon_event import CouponEvent
from structured_instruments.models.instrument import Instrument, InstrumentUnderlying
from structured_instruments.models.position import PortfolioPosition
from structured_instruments.models.price import DailyPrice
from structured_instruments.models.underlying import Underlying
from structured_instruments.models.valuation_date import ValuationDate

router = APIRouter(prefix="/api/export", tags=["export"])

# Shared styles
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill(start_color="374151", end_color="374151", fill_type="solid")
HEADER_ALIGNMENT = Alignment(horizontal="center", wrap_text=True)


def _style_header_row(ws, num_cols: int) -> None:
    """Apply consistent header styling to the first row of a worksheet."""
    for col_idx in range(1, num_cols + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT


def _auto_column_widths(ws) -> None:
    """Set column widths based on content."""
    for column_cells in ws.columns:
        max_length = 0
        column_letter = column_cells[0].column_letter
        for cell in column_cells:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[column_letter].width = min(max_length + 4, 40)


def _build_portfolio_summary_sheet(wb: Workbook, session: Session, account_ids: list[int]) -> None:
    """Sheet 1: Portfolio Summary."""
    ws = wb.active
    ws.title = "Portfolio Summary"

    # Gather data for selected accounts
    principal_stmt = select(
        func.coalesce(func.sum(PortfolioPosition.principal), 0.0)
    ).where(PortfolioPosition.account_id.in_(account_ids))
    total_principal = float(session.execute(principal_stmt).scalar_one())

    # Get positions for coupon income calculation
    pos_stmt = (
        select(PortfolioPosition)
        .join(Instrument)
        .where(PortfolioPosition.account_id.in_(account_ids))
    )
    positions = list(session.execute(pos_stmt).scalars().all())
    user_instrument_ids = list({p.instrument_id for p in positions})

    total_income = 0.0
    for pos in positions:
        inst = pos.instrument
        paid_count_stmt = select(func.count()).where(
            CouponEvent.instrument_id == inst.id,
            CouponEvent.is_paid.is_(True),
        )
        paid_count = session.execute(paid_count_stmt).scalar_one()
        total_income += paid_count * pos.principal * inst.coupon_rate_monthly

    total_instruments = len(user_instrument_ids)

    active_count = 0
    called_count = 0
    if user_instrument_ids:
        active_count = session.execute(
            select(func.count()).where(
                Instrument.id.in_(user_instrument_ids),
                Instrument.status == "active",
            )
        ).scalar_one()
        called_count = session.execute(
            select(func.count()).where(
                Instrument.id.in_(user_instrument_ids),
                Instrument.status == "called",
            )
        ).scalar_one()

    # Holding period: earliest purchase_date to today
    earliest_stmt = select(func.min(PortfolioPosition.purchase_date)).where(
        PortfolioPosition.account_id.in_(account_ids)
    )
    earliest_date = session.execute(earliest_stmt).scalar_one_or_none()
    holding_days = (date.today() - earliest_date).days if earliest_date else 0

    # Write summary as label-value pairs
    headers = ["Metric", "Value"]
    ws.append(headers)
    _style_header_row(ws, len(headers))

    rows = [
        ("Total Principal", f"${total_principal:,.2f}"),
        ("Total Coupon Income", f"${total_income:,.2f}"),
        ("Total Instruments", total_instruments),
        ("Active Instruments", active_count),
        ("Called Instruments", called_count),
        ("Holding Period (days)", holding_days),
        ("Export Date", date.today().isoformat()),
    ]
    for label, value in rows:
        ws.append([label, value])

    _auto_column_widths(ws)


def _build_instruments_sheet(wb: Workbook, session: Session, account_ids: list[int]) -> None:
    """Sheet 2: Instruments (with position data from selected accounts)."""
    ws = wb.create_sheet("Instruments")

    headers = [
        "CUSIP", "Issuer", "Product Type", "Annual Coupon Rate",
        "Monthly Coupon Rate", "Coupon Barrier", "Downside Barrier",
        "Autocall Barrier", "Pricing Date", "Maturity Date", "Status",
        "Principal", "Num Notes",
    ]
    ws.append(headers)
    _style_header_row(ws, len(headers))

    stmt = (
        select(Instrument)
        .join(PortfolioPosition)
        .where(PortfolioPosition.account_id.in_(account_ids))
        .distinct()
        .order_by(Instrument.cusip)
    )
    instruments = list(session.execute(stmt).scalars().all())

    for inst in instruments:
        # Sum principal/notes across selected accounts
        pos_stmt = select(PortfolioPosition).where(
            PortfolioPosition.instrument_id == inst.id,
            PortfolioPosition.account_id.in_(account_ids),
        )
        positions = list(session.execute(pos_stmt).scalars().all())
        total_principal = sum(p.principal for p in positions) if positions else 0
        total_notes = sum(p.num_notes for p in positions) if positions else 0
        ws.append([
            inst.cusip,
            inst.issuer,
            inst.product_type or "",
            f"{inst.coupon_rate_annual * 100:.2f}%",
            f"{inst.coupon_rate_monthly * 100:.4f}%",
            f"{inst.coupon_barrier_pct * 100:.1f}%",
            f"{inst.downside_barrier_pct * 100:.1f}%",
            f"{inst.autocall_barrier_pct * 100:.1f}%",
            inst.pricing_date.isoformat(),
            inst.maturity_date.isoformat(),
            inst.status,
            f"${total_principal:,.2f}" if positions else "",
            total_notes if positions else "",
        ])

    _auto_column_widths(ws)


def _build_underlyings_sheet(wb: Workbook, session: Session, account_ids: list[int]) -> None:
    """Sheet 3: Underlyings (for selected accounts' instruments)."""
    ws = wb.create_sheet("Underlyings")

    headers = [
        "Ticker", "Name", "CUSIP", "Initial Price", "Current Price",
        "% of Initial", "Coupon Barrier Price", "Downside Barrier Price",
    ]
    ws.append(headers)
    _style_header_row(ws, len(headers))

    today = date.today()

    # Get instrument IDs from selected accounts
    user_inst_ids = [
        row[0] for row in session.execute(
            select(PortfolioPosition.instrument_id).where(
                PortfolioPosition.account_id.in_(account_ids)
            ).distinct()
        ).all()
    ]

    if not user_inst_ids:
        _auto_column_widths(ws)
        return

    stmt = (
        select(InstrumentUnderlying)
        .join(Instrument)
        .join(Underlying, InstrumentUnderlying.underlying_id == Underlying.id)
        .where(InstrumentUnderlying.instrument_id.in_(user_inst_ids))
        .order_by(Instrument.cusip, Underlying.ticker)
    )
    ius = list(session.execute(stmt).scalars().all())

    for iu in ius:
        und = iu.underlying
        inst = iu.instrument

        # Latest price
        price_stmt = (
            select(DailyPrice.close)
            .where(
                DailyPrice.underlying_id == und.id,
                DailyPrice.date <= today,
            )
            .order_by(DailyPrice.date.desc())
            .limit(1)
        )
        current_price = session.execute(price_stmt).scalar_one_or_none()
        pct_of_initial = (current_price / iu.initial_price) if current_price else None

        ws.append([
            und.ticker,
            und.name or "",
            inst.cusip,
            round(iu.initial_price, 2),
            round(current_price, 2) if current_price else "",
            f"{pct_of_initial * 100:.2f}%" if pct_of_initial else "",
            round(iu.coupon_barrier_price, 2),
            round(iu.downside_barrier_price, 2),
        ])

    _auto_column_widths(ws)


def _build_coupon_history_sheet(wb: Workbook, session: Session, account_ids: list[int]) -> None:
    """Sheet 4: Coupon History (amounts computed per selected accounts)."""
    ws = wb.create_sheet("Coupon History")

    headers = [
        "Date", "CUSIP", "Status", "Amount",
        "Worst Performer Ticker", "Worst Performer %",
    ]
    ws.append(headers)
    _style_header_row(ws, len(headers))

    # Get positions for coupon amount calculation
    pos_stmt = select(PortfolioPosition).where(PortfolioPosition.account_id.in_(account_ids))
    positions = list(session.execute(pos_stmt).scalars().all())
    # Sum principal per instrument
    principal_by_inst: dict[int, float] = {}
    for p in positions:
        principal_by_inst[p.instrument_id] = principal_by_inst.get(p.instrument_id, 0) + p.principal
    user_inst_ids = list(principal_by_inst.keys())

    if not user_inst_ids:
        _auto_column_widths(ws)
        return

    stmt = (
        select(CouponEvent)
        .join(Instrument)
        .where(CouponEvent.instrument_id.in_(user_inst_ids))
        .order_by(CouponEvent.evaluation_date.desc(), Instrument.cusip)
    )
    events = list(session.execute(stmt).scalars().all())

    for ev in events:
        # Compute coupon amount using summed principal
        principal = principal_by_inst.get(ev.instrument_id, 0)
        if ev.is_paid and principal > 0:
            amount = principal * ev.instrument.coupon_rate_monthly
            amount_str = f"${amount:,.2f}"
        else:
            amount_str = ""

        ws.append([
            ev.evaluation_date.isoformat(),
            ev.instrument.cusip,
            "Paid" if ev.is_paid else "Missed",
            amount_str,
            ev.worst_performer_ticker or "",
            f"{ev.worst_performer_pct * 100:.2f}%" if ev.worst_performer_pct is not None else "",
        ])

    _auto_column_widths(ws)


def _build_upcoming_dates_sheet(wb: Workbook, session: Session, account_ids: list[int]) -> None:
    """Sheet 5: Upcoming Dates (for selected accounts' instruments)."""
    ws = wb.create_sheet("Upcoming Dates")

    headers = [
        "Date", "CUSIP", "Issuer", "Sequence Number", "Autocall Eligible",
    ]
    ws.append(headers)
    _style_header_row(ws, len(headers))

    user_inst_ids = [
        row[0] for row in session.execute(
            select(PortfolioPosition.instrument_id).where(
                PortfolioPosition.account_id.in_(account_ids)
            ).distinct()
        ).all()
    ]

    if not user_inst_ids:
        _auto_column_widths(ws)
        return

    today = date.today()
    stmt = (
        select(ValuationDate)
        .join(Instrument)
        .where(
            ValuationDate.date > today,
            ValuationDate.instrument_id.in_(user_inst_ids),
        )
        .order_by(ValuationDate.date, Instrument.cusip)
    )
    vds = list(session.execute(stmt).scalars().all())

    for vd in vds:
        inst = vd.instrument
        ws.append([
            vd.date.isoformat(),
            inst.cusip,
            inst.issuer,
            vd.sequence_number,
            "Yes" if vd.autocall_eligible else "No",
        ])

    _auto_column_widths(ws)


@router.get("/excel")
def export_excel(
    account_ids: list[int] = Depends(get_account_filter),
    session: Session = Depends(get_db),
):
    """Generate and return an Excel workbook with the selected accounts' portfolio data."""
    wb = Workbook()

    _build_portfolio_summary_sheet(wb, session, account_ids)
    _build_instruments_sheet(wb, session, account_ids)
    _build_underlyings_sheet(wb, session, account_ids)
    _build_coupon_history_sheet(wb, session, account_ids)
    _build_upcoming_dates_sheet(wb, session, account_ids)

    # Serialize to bytes
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"structured_instruments_{date.today().isoformat()}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
